from io import BytesIO
from html import escape
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

def format_html_preview(text):
    blocks=[]

    for line in text.splitlines():
        line=line.strip()

        if not line:
            continue

        if line.isupper() or line.startswith("#"):
            blocks.append(
                f"<h3>{escape(line.lstrip('#').strip())}</h3>"
            )
        else:
            blocks.append(f"<p>{escape(line)}</p>")

    return "".join(blocks)

def format_docx(text,doc_type):
    doc=Document()

    section=doc.sections[0]
    section.top_margin=section.bottom_margin=0.7
    section.left_margin=section.right_margin=0.8

    title=doc.add_paragraph()
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER

    run=title.add_run(doc_type.upper())
    run.bold=True
    run.font.size=Pt(16)

    for line in text.splitlines():
        line=line.strip()

        if not line:
            continue

        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(6)

        r=p.add_run(line)
        r.font.name="Times New Roman"
        r.font.size=Pt(11)

    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER

    footer_run=footer.add_run(
        "Generated with LegalEase | Educational use only"
    )
    footer_run.font.size=Pt(9)

    output=BytesIO()
    doc.save(output)
    output.seek(0)

    return output.getvalue()

class LegalPDF(FPDF):

    def header(self):
        self.set_font("Helvetica","B",13)
        self.cell(0,10,"LegalEase",align="C")
        self.ln(8)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica","",8)
        self.cell(
            0,
            10,
            "Generated with LegalEase | Educational use only",
            align="C"
        )

def format_pdf(text,doc_type):
    pdf=LegalPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=18
    )

    pdf.add_page()

    pdf.set_font("Helvetica","B",15)
    pdf.multi_cell(
        0,
        9,
        doc_type,
        align="C"
    )

    pdf.ln(5)
    pdf.set_font("Helvetica","",11)

    for line in text.splitlines():
        line=line.strip()

        if not line:
            pdf.ln(3)
            continue

        safe=line.encode(
            "latin-1",
            "replace"
        ).decode("latin-1")

        pdf.multi_cell(0,7,safe)
        pdf.ln(1)

    return bytes(pdf.output())